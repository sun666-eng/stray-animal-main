from pathlib import Path
import re
from docx import Document

path = Path(r"D:\2026.4.7\Myproject\stray-animal-main\outputs\食品质量与安全542403040115覃祝_论文3300字完成版.docx")
doc = Document(path)
texts = [p.text.strip() for p in doc.paragraphs]
title_idx = next(i for i, x in enumerate(texts) if x == "题目：食品安全快速检测技术研究综述")
ref_idx = next(i for i, x in enumerate(texts) if i > title_idx and x == "参考文献")
body = "".join(texts[title_idx + 1:ref_idx])
chinese = len(re.findall(r"[\u4e00-\u9fff]", body))

assert chinese >= 3300, chinese
assert len(doc.inline_shapes) == 2
assert len(doc.tables) == 3
assert any("智能手机、微流控与人工智能融合" in x for x in texts)
assert any("Intelligent biosensing strategies" in x for x in texts)

print(f"file={path}")
print(f"paragraphs={len(texts)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
print(f"paper_total_chars={len(body)}")
print(f"paper_chinese_chars={chinese}")
print("validation=PASS")
