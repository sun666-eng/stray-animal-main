from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\2026.4.7\Myproject\stray-animal-main")
SRC = ROOT / "流浪动物救助管理系统-实验报告-含图1-1.docx"
if not SRC.exists():
    SRC = ROOT / "流浪动物救助管理系统-实验报告.docx"
OUT = ROOT / "流浪动物救助管理系统-实验报告-最终版.docx"
FIG_DIR = ROOT / "figures"


def fnt(size, bold=False):
    return ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc", size)


def wrap_text(draw, text, font, max_width):
    lines = []
    for raw in text.split("\n"):
        line = ""
        for ch in raw:
            if draw.textbbox((0, 0), line + ch, font=font)[2] <= max_width:
                line += ch
            else:
                if line:
                    lines.append(line)
                line = ch
        if line:
            lines.append(line)
    return lines or [""]


def text_center(draw, box, text, font, fill="#172033", gap=7):
    lines = wrap_text(draw, text, font, box[2] - box[0] - 18)
    hs, ws = [], []
    for line in lines:
        b = draw.textbbox((0, 0), line, font=font)
        ws.append(b[2] - b[0])
        hs.append(b[3] - b[1])
    total_h = sum(hs) + gap * (len(lines) - 1)
    y = box[1] + ((box[3] - box[1]) - total_h) / 2
    for i, line in enumerate(lines):
        x = box[0] + ((box[2] - box[0]) - ws[i]) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += hs[i] + gap


def text_left(draw, xy, text, font, fill="#172033", max_width=360, gap=6):
    x, y = xy
    for line in wrap_text(draw, text, font, max_width):
        draw.text((x, y), line, font=font, fill=fill)
        y += draw.textbbox((0, 0), line, font=font)[3] + gap


def box(draw, rect, text, fill="#FFFFFF", outline="#334155", width=2, radius=18, font=None, color="#172033"):
    draw.rounded_rectangle(rect, radius=radius, fill=fill, outline=outline, width=width)
    text_center(draw, rect, text, font or fnt(24), color)


def line_arrow(draw, start, end, color="#475569", width=3):
    import math
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 13
    p1 = (end[0] - size * math.cos(angle - math.pi / 6), end[1] - size * math.sin(angle - math.pi / 6))
    p2 = (end[0] - size * math.cos(angle + math.pi / 6), end[1] - size * math.sin(angle + math.pi / 6))
    draw.polygon([end, p1, p2], fill=color)


def title(draw, text):
    draw.text((70, 42), text, font=fnt(42, True), fill="#0F172A")


def canvas(title_text, size=(1800, 1180)):
    img = Image.new("RGB", size, "#F8FAFC")
    d = ImageDraw.Draw(img)
    title(d, title_text)
    return img, d


def save(img, name):
    FIG_DIR.mkdir(exist_ok=True)
    path = FIG_DIR / name
    img.save(path, quality=96)
    return path


def draw_actor(draw, x, y, label, color):
    draw.ellipse((x - 22, y - 62, x + 22, y - 18), fill=color)
    draw.line((x, y - 18, x, y + 52), fill=color, width=6)
    draw.line((x - 48, y + 4, x + 48, y + 4), fill=color, width=6)
    draw.line((x, y + 52, x - 42, y + 112), fill=color, width=6)
    draw.line((x, y + 52, x + 42, y + 112), fill=color, width=6)
    text_center(draw, (x - 105, y + 122, x + 105, y + 176), label, fnt(23, True))


def use_case_image(name, title_text):
    img, d = canvas(title_text)
    sys = (430, 145, 1370, 1040)
    d.rounded_rectangle(sys, radius=34, fill="#EFF6FF", outline="#2563EB", width=4)
    text_center(d, (sys[0], 158, sys[2], 220), "系统边界：流浪动物救助管理系统", fnt(28, True), "#1E3A8A")
    actors = [
        (180, 260, "访客", "#0284C7"),
        (180, 700, "普通用户", "#059669"),
        (1600, 285, "工作人员", "#EA580C"),
        (1600, 720, "管理员", "#DC2626"),
    ]
    for x, y, label, color in actors:
        draw_actor(d, x, y, label, color)
    cases = [
        ("浏览动物", 610, 310), ("查看公告", 610, 455), ("注册登录", 610, 600),
        ("提交领养申请", 830, 370), ("查看个人申请", 830, 540), ("提交救助咨询", 830, 710),
        ("动物管理", 1090, 310), ("领养/义工审核", 1090, 455), ("救助与回访处理", 1090, 600),
        ("用户角色权限", 1090, 760), ("资金公告管理", 820, 865), ("Excel 数据导出", 610, 790),
    ]
    for text, x, y in cases:
        d.ellipse((x - 135, y - 48, x + 135, y + 48), fill="#FFFFFF", outline="#60A5FA", width=3)
        text_center(d, (x - 135, y - 48, x + 135, y + 48), text, fnt(21))
    for start, ends, color in [
        ((285, 360), [(475, 310), (475, 455)], "#0284C7"),
        ((285, 790), [(475, 600), (695, 370), (695, 540), (695, 710)], "#059669"),
        ((1495, 380), [(1225, 310), (1225, 455), (1225, 600)], "#EA580C"),
        ((1495, 805), [(1225, 760), (955, 865), (745, 790)], "#DC2626"),
    ]:
        for end in ends:
            d.line([start, end], fill=color, width=3)
    return save(img, name)


def flowchart(name, title_text, steps, colors=None):
    img, d = canvas(title_text, (1800, 1320))
    colors = colors or ["#DBEAFE"] * len(steps)
    x = 220
    w = 1360
    y = 150
    h = 82
    for i, step in enumerate(steps):
        rect = (x, y, x + w, y + h)
        if step.startswith("?"):
            cx, cy = x + w // 2, y + h // 2
            pts = [(cx, y - 8), (x + w - 60, cy), (cx, y + h + 8), (x + 60, cy)]
            d.polygon(pts, fill=colors[i], outline="#334155")
            text_center(d, rect, step[1:], fnt(22, True))
        else:
            box(d, rect, step, colors[i], "#334155", 2, 18, fnt(22))
        if i < len(steps) - 1:
            line_arrow(d, (x + w // 2, y + h), (x + w // 2, y + 128))
        y += 145
    return save(img, name)


def sequence(name, title_text, participants, messages):
    img, d = canvas(title_text, (1800, 1180))
    xs = [180 + i * (1440 // (len(participants) - 1)) for i in range(len(participants))]
    y0, y1 = 150, 1040
    for x, p in zip(xs, participants):
        box(d, (x - 115, y0, x + 115, y0 + 62), p, "#E0F2FE", "#0284C7", 2, 14, fnt(20, True))
        d.line((x, y0 + 62, x, y1), fill="#94A3B8", width=2)
    y = 250
    for idx, (a, b, txt) in enumerate(messages, 1):
        x1, x2 = xs[a], xs[b]
        line_arrow(d, (x1, y), (x2, y), "#475569", 3)
        mid = (x1 + x2) / 2
        d.text((mid - 135, y - 30), f"{idx}. {txt}", font=fnt(18), fill="#111827")
        y += 78
    return save(img, name)


def architecture():
    img, d = canvas("图4-1 系统总体架构图")
    layers = [
        ("表现层", "前台页面：动物浏览、公告、领养申请、义工申请、救助咨询\n后台页面：用户、角色、权限、动物、领养、义工、救助、资金、公告", "#DBEAFE"),
        ("接口层", "Spring MVC Controller：统一 REST API、Result 返回、参数接收、文件上传、Excel 导出", "#E0F2FE"),
        ("业务层", "Service：用户登录、权限补齐、本人数据校验、领养审核、救助处理、导出数据组装", "#ECFDF5"),
        ("持久层", "MyBatis-Plus Mapper 与 XML：分页查询、组合主键、实体映射、数据库访问", "#FFF7ED"),
        ("数据与资源层", "MySQL 数据库、upload 文件目录、日志文件、repair_permissions.sql 修复脚本", "#FEF2F2"),
    ]
    y = 150
    for label, desc, fill in layers:
        box(d, (120, y, 420, y + 110), label, fill, "#334155", 3, 18, fnt(28, True))
        box(d, (470, y, 1680, y + 110), desc, "#FFFFFF", "#CBD5E1", 2, 18, fnt(22))
        if y < 150 + 4 * 165:
            line_arrow(d, (900, y + 110), (900, y + 160), "#64748B")
        y += 165
    return save(img, "图4-1-系统总体架构图.png")


def er_diagram(name, title_text):
    img, d = canvas(title_text, (1900, 1350))
    entities = {
        "t_user\nid PK\nusername\npassword\nrole(JSON)": (80, 160),
        "t_role\nid PK\nname\npermission(JSON)": (450, 160),
        "t_permission\nid PK\npath\nflag": (820, 160),
        "t_animal\nid PK\ntname\ntstate": (80, 520),
        "t_adopt\naid PK/FK\nuid PK/FK\nvstate": (450, 520),
        "t_proof\nid PK\npaid\npuid": (820, 520),
        "t_visit\nid PK\npet_id\nuid\nstate": (1190, 520),
        "t_volunteer\nid PK\nname\ntel\nvstate": (80, 900),
        "t_help\nid PK\nuid\ntitle\nstatus": (450, 900),
        "t_notice\nid PK\ntitle\ncontent": (820, 900),
        "t_account\nid PK\nalabel\navalue": (1190, 900),
    }
    centers = {}
    for text, (x, y) in entities.items():
        rect = (x, y, x + 280, y + 220)
        box(d, rect, text, "#FFFFFF", "#334155", 2, 16, fnt(20))
        header = (x, y, x + 280, y + 48)
        d.rounded_rectangle(header, radius=16, fill="#DBEAFE", outline="#334155", width=0)
        text_center(d, header, text.split("\n")[0], fnt(21, True), "#1E3A8A")
        centers[text.split("\n")[0]] = (x + 140, y + 110)
    def rel(a, b, label):
        ax, ay = centers[a]; bx, by = centers[b]
        d.line((ax, ay, bx, by), fill="#475569", width=3)
        mx, my = (ax + bx) // 2, (ay + by) // 2
        box(d, (mx - 70, my - 22, mx + 70, my + 22), label, "#F8FAFC", "#94A3B8", 1, 12, fnt(16))
    rel("t_user", "t_role", "拥有角色")
    rel("t_role", "t_permission", "包含权限")
    rel("t_user", "t_adopt", "提交")
    rel("t_animal", "t_adopt", "被申请")
    rel("t_adopt", "t_proof", "上传凭证")
    rel("t_animal", "t_visit", "回访")
    rel("t_user", "t_help", "发起救助")
    rel("t_user", "t_volunteer", "申请义工")
    return save(img, name)


def class_diagram(name, title_text):
    img, d = canvas(title_text, (1900, 1250))
    classes = [
        ("UserController\n+ login()\n+ register()\n+ findPage()", 80, 170, "#DBEAFE"),
        ("UserService\n+ login()\n+ register()\n+ fillPermissions()", 450, 170, "#E0F2FE"),
        ("UserMapper\n+ BaseMapper<User>", 820, 170, "#F0FDFA"),
        ("User\n- id\n- username\n- role\n- permission", 1190, 170, "#FFFFFF"),
        ("AuthInterceptor\n+ preHandle()\n- hasApiPermission()\n- getCurrentUser()", 80, 520, "#FEF3C7"),
        ("PermissionUtil\n+ hasFlag()\n+ hasAnyFlag()", 450, 520, "#FEF3C7"),
        ("JwtUtil\n+ createToken()\n+ validate()\n+ getUserId()", 820, 520, "#FEF3C7"),
        ("ExcelExportUtil\n+ export()", 1190, 520, "#FEF3C7"),
        ("Animal/Adopt/Help\nVolunteer/Visit/Proof\nAccount/Notice", 450, 860, "#FFFFFF"),
        ("各模块 Controller\nAnimalController\nAdoptController\nHelpController\n...", 80, 860, "#DBEAFE"),
        ("各模块 Service\n继承 ServiceImpl", 820, 860, "#E0F2FE"),
        ("各模块 Mapper\n继承 BaseMapper", 1190, 860, "#F0FDFA"),
    ]
    centers = []
    for text, x, y, fill in classes:
        box(d, (x, y, x + 310, y + 210), text, fill, "#334155", 2, 16, fnt(19))
        centers.append((x + 155, y + 105))
    for a, b in [(0, 1), (1, 2), (2, 3), (4, 1), (4, 5), (4, 6), (9, 10), (10, 11), (10, 8), (7, 9)]:
        line_arrow(d, centers[a], centers[b], "#475569", 2)
    return save(img, name)


def permission_flow():
    steps = [
        "浏览器发送请求，前端携带 Authorization: Bearer token",
        "?请求是否为登录、注册、公告、动物浏览或文件下载等公开接口",
        "公开接口直接放行，非公开接口进入身份识别",
        "解析 JWT；若无有效 token，则读取 Session 中用户信息",
        "?是否取得当前用户",
        "未登录返回 401，并由前端跳转登录页",
        "根据用户角色补齐 permission.flag 权限集合",
        "?请求是否为“我的”接口",
        "校验 uid、username、phone、email 是否属于当前用户",
        "后台管理接口按模块 flag 判断，例如 user、role、animal、adopt",
        "?是否具备对应权限",
        "放行请求，Controller 执行业务逻辑",
        "无权限返回 403，页面提示无权访问",
    ]
    return flowchart("图6-2-权限校验流程图.png", "图6-2 权限校验流程图", steps,
                     ["#DBEAFE", "#FEF3C7", "#E0F2FE", "#E0F2FE", "#FEF3C7", "#FEE2E2",
                      "#ECFDF5", "#FEF3C7", "#ECFDF5", "#DBEAFE", "#FEF3C7", "#DCFCE7", "#FEE2E2"])


def build_all_figures():
    paths = {}
    paths["图2-1 系统总体用例图"] = use_case_image("图2-1-系统总体用例图.png", "图2-1 系统总体用例图")
    paths["图3-1 领养申请业务活动图"] = flowchart(
        "图3-1-领养申请业务活动图.png", "图3-1 领养申请业务活动图",
        ["用户浏览待领养动物", "进入动物详情并确认领养意向", "?是否已经登录", "未登录则跳转登录页",
         "填写领养申请表并提交", "系统保存 t_adopt 待审核记录", "管理员查看申请资料",
         "?审核是否通过", "通过：更新申请状态并允许上传凭证", "未通过：记录结果并反馈用户", "用户在我的领养中查看状态"],
        ["#DBEAFE", "#DBEAFE", "#FEF3C7", "#FEE2E2", "#E0F2FE", "#ECFDF5", "#FFF7ED",
         "#FEF3C7", "#DCFCE7", "#FEE2E2", "#DBEAFE"])
    paths["图3-2 义工申请与审核流程图"] = flowchart(
        "图3-2-义工申请与审核流程图.png", "图3-2 义工申请与审核流程图",
        ["用户进入义工申请页面", "填写姓名、电话、邮箱、所在地、空闲时间和能力说明", "提交申请并写入 t_volunteer",
         "系统仅允许用户查看自己的义工申请", "工作人员分页查看义工申请", "?是否符合义工要求",
         "通过：更新 vstate 为通过", "驳回：更新 vstate 并提示用户", "用户在我的义工申请中查看审核结果"],
        ["#DBEAFE", "#E0F2FE", "#ECFDF5", "#DBEAFE", "#FFF7ED", "#FEF3C7", "#DCFCE7", "#FEE2E2", "#DBEAFE"])
    paths["图3-3 救助咨询顺序图"] = sequence(
        "图3-3-救助咨询顺序图.png", "图3-3 救助咨询顺序图",
        ["用户页面", "AuthInterceptor", "HelpController", "HelpService", "HelpMapper", "MySQL"],
        [(0, 1, "提交救助请求"), (1, 2, "校验 token 与 im/help 权限"), (2, 3, "组装 Help 实体"),
         (3, 4, "调用保存方法"), (4, 5, "写入 t_help"), (5, 4, "返回保存结果"),
         (4, 3, "返回业务结果"), (3, 0, "Result.success")])
    paths["图3-4 登录认证与权限校验顺序图"] = sequence(
        "图3-4-登录认证与权限校验顺序图.png", "图3-4 登录认证与权限校验顺序图",
        ["登录页", "UserController", "UserService", "UserMapper", "JwtUtil", "Session"],
        [(0, 1, "POST /api/user/login"), (1, 2, "调用 login"), (2, 3, "查询用户"),
         (3, 2, "返回用户记录"), (2, 2, "校验/升级 BCrypt 密码"), (2, 2, "补齐角色权限"),
         (1, 4, "创建 JWT"), (1, 5, "写入 Session"), (1, 0, "返回 token 与用户权限")])
    paths["图4-1 系统总体架构图"] = architecture()
    paths["图5-1 数据库 ER 图"] = er_diagram("图5-1-数据库ER图.png", "图5-1 数据库 ER 图")
    paths["图6-1 系统核心类图"] = class_diagram("图6-1-系统核心类图.png", "图6-1 系统核心类图")
    paths["图6-2 权限校验流程图"] = permission_flow()
    paths["图9-1 用例图粘贴位置"] = paths["图2-1 系统总体用例图"]
    paths["图9-2 活动图粘贴位置"] = paths["图3-1 领养申请业务活动图"]
    paths["图9-3 类图粘贴位置"] = paths["图6-1 系统核心类图"]
    paths["图9-4 顺序图粘贴位置"] = paths["图3-4 登录认证与权限校验顺序图"]
    paths["图9-5 ER 图粘贴位置"] = paths["图5-1 数据库 ER 图"]
    return paths


def set_run_font(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def replace_paragraph_text(p, text):
    for r in p.runs:
        r.text = ""
    r = p.add_run(text)
    set_run_font(r, 10.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None


def insert_images_and_toc_marker(paths):
    shutil.copyfile(SRC, OUT)
    doc = Document(str(OUT))
    caption_replacements = {
        "图9-1 用例图粘贴位置": "图9-1 系统用例图补充说明",
        "图9-2 活动图粘贴位置": "图9-2 业务活动图补充说明",
        "图9-3 类图粘贴位置": "图9-3 系统类图补充说明",
        "图9-4 顺序图粘贴位置": "图9-4 登录认证顺序图补充说明",
        "图9-5 ER 图粘贴位置": "图9-5 数据库 ER 图补充说明",
    }
    for p in doc.paragraphs:
        if p.text in caption_replacements:
            replace_paragraph_text(p, caption_replacements[p.text])

    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith("【此处粘贴：") and text.endswith("】"):
            key = text.replace("【此处粘贴：", "").replace("】", "")
            img_path = paths.get(key)
            if img_path:
                p.clear()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = None
                p.add_run().add_picture(str(img_path), width=Inches(6.35))

    if not any("[[TOC_HERE]]" in p.text for p in doc.paragraphs):
        first_heading = next(p for p in doc.paragraphs if p.text.strip().startswith("1  引言"))
        toc_title = first_heading.insert_paragraph_before("目录")
        toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        toc_title.paragraph_format.first_line_indent = None
        for r in toc_title.runs:
            set_run_font(r, 16, True)
        toc_marker = first_heading.insert_paragraph_before("[[TOC_HERE]]")
        toc_marker.paragraph_format.first_line_indent = None
        toc_break = first_heading.insert_paragraph_before("")
        toc_break.add_run().add_break(WD_BREAK.PAGE)
    doc.save(str(OUT))
    print(OUT)


def main():
    paths = build_all_figures()
    insert_images_and_toc_marker(paths)
    print("figures", len(paths))


if __name__ == "__main__":
    main()
