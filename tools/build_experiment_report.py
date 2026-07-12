from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\2026.4.7\Myproject\stray-animal-main")
TEMPLATE = Path(r"D:\迅雷\实验报告模板(2).docx")
OUT = ROOT / "流浪动物救助管理系统-实验报告.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def clear_body_after_cover(doc):
    body = doc._body._element
    keep = 18
    paras_seen = 0
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:p"):
            paras_seen += 1
            if paras_seen <= keep:
                continue
        body.remove(child)


def set_run_font(run, size=12, bold=False, color=None):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [
        ("Heading 1", 16, "1F4E79"),
        ("Heading 2", 14, "1F4E79"),
        ("Heading 3", 12, "1F4E79"),
    ]:
        if name not in doc.styles:
            continue
        s = doc.styles[name]
        s.font.name = "宋体"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        s._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.first_line_indent = None
        s.paragraph_format.space_before = Pt(10)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.line_spacing = 1.5


def replace_cover_text(doc):
    replacements = {
        "项目名称：                           ": "项目名称：流浪动物救助管理系统",
        "组名：                               ": "组名：第 1 组（请手动填写）",
        "小组成员：                           ": "小组成员：（请手动填写姓名与学号）",
        "指导教师：                           ": "指导教师：（请手动填写）",
    }
    for p in doc.paragraphs:
        text = p.text
        if text in replacements:
            for r in p.runs:
                r.text = ""
            run = p.add_run(replacements[text])
            set_run_font(run, 12)


def add_para(doc, text, style=None, indent=True):
    p = doc.add_paragraph(style=style)
    if not indent:
        p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if indent and style is None:
        p.paragraph_format.first_line_indent = Pt(24)
    run = p.add_run(text)
    set_run_font(run, 12)
    return p


def add_heading(doc, text, level):
    if level == 1:
        count = getattr(add_heading, "h1_count", 0)
        if count > 0:
            doc.add_page_break()
        setattr(add_heading, "h1_count", count + 1)
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, 16 if level == 1 else 14 if level == 2 else 12, True, "1F4E79")
    return p


def add_bullets(doc, items):
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run("（{}）{}".format(idx, item))
        set_run_font(run, 12)


def add_numbered(doc, items):
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run("{}. {}".format(idx, item))
        set_run_font(run, 12)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "E8EEF5")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(hdr[i], widths[i])
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, 10.5, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cells[i], widths[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.first_line_indent = None
                p.paragraph_format.line_spacing = 1.2
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_run_font(r, 10.5)
    doc.add_paragraph()
    return table


def add_figure_placeholder(doc, caption):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("【此处粘贴：" + caption + "】")
    set_run_font(run, 12, True, "7A5A00")
    p2 = doc.add_paragraph(caption)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.first_line_indent = None
    for r in p2.runs:
        set_run_font(r, 10.5)


def build_report(doc):
    add_heading(doc, "1  引言", 1)
    add_heading(doc, "1.1 系统背景", 2)
    add_para(doc, "随着城市化进程不断加快，社区、高校、商业街区等公共区域中的流浪动物数量持续增加。流浪动物救助工作通常涉及动物信息登记、照片上传、领养申请审核、志愿者招募、救助线索跟进、领养后回访、资金收支公示等多个环节。如果这些数据长期依靠纸质表格、即时通信群或人工台账维护，容易出现信息分散、状态不一致、审核过程不可追溯、公众查询不方便等问题。")
    add_para(doc, "本实验以“流浪动物救助管理系统”为研究对象，围绕系统分析与建模课程中需求分析、用例建模、业务流程建模、静态结构建模、动态交互建模和数据库建模等知识点，对一个具有实际业务背景的 Web 管理系统进行分析、设计、实现和测试。系统采用 B/S 架构，前端以 HTML、Vue、jQuery、Element UI 等技术组织页面交互，后端采用 Spring Boot、MyBatis-Plus、MySQL、JWT、拦截器和 Apache POI 等技术完成接口、数据访问、安全控制和 Excel 导出。")
    add_para(doc, "系统服务对象包括普通用户、志愿者和管理员。普通用户可以浏览待领养动物、查看公告、提交领养申请、查看个人申请记录、提交义工申请、提交救助咨询和查看自己的救助记录；志愿者或工作人员可以协助维护动物资料、领养审核、义工审核、回访记录和救助处理；管理员负责用户、角色、权限、公告、资金、动物、领养、凭证、回访、义工和救助等后台数据的统一管理。")
    add_heading(doc, "1.2 实验目的与意义", 2)
    add_para(doc, "本实验的目的不是单纯完成一个 CRUD 系统，而是通过一个包含前台用户流程和后台管理流程的综合案例，训练对真实需求进行抽象、分解和验证的能力。通过本系统，可以把课程中的模型图与工程代码对应起来：用例图对应角色与功能边界，活动图对应业务状态流转，类图对应实体对象与服务对象，顺序图对应一次请求从页面到控制器、服务层、数据层再返回页面的调用过程，ER 图对应数据库表之间的关系。")
    add_para(doc, "从应用价值看，系统可以提升救助组织的信息化水平，降低人工记录成本，提高领养审核透明度，并通过公告、资金公示和回访记录增强社会公众对救助工作的信任。从课程实践看，系统具备较完整的业务链路，既有公开浏览场景，也有登录后个人业务场景，还有严格权限控制下的后台管理场景，适合用于展示面向对象分析、分层架构设计和权限模型设计。")
    add_heading(doc, "1.3 报告范围", 2)
    add_para(doc, "本报告围绕当前项目代码和数据库脚本展开，内容包括系统需求分析、功能模块划分、非功能需求、总体架构设计、数据库设计、核心业务流程、关键类与接口设计、实现说明、测试方案和实验总结。报告中需要绘制的用例图、活动图、类图、顺序图、ER 图等位置均保留图号和占位说明，后续可根据最终绘图结果手动粘贴。")
    add_figure_placeholder(doc, "图1-1 流浪动物救助管理系统应用场景示意图")

    add_heading(doc, "2  系统需求分析", 1)
    add_heading(doc, "2.1 用户角色分析", 2)
    add_para(doc, "系统中的参与者可分为三类。第一类是未登录访客，主要进行动物浏览、动物详情查看、公告列表查看和公告详情查看。第二类是普通用户，登录后可以提交领养申请、查看自己的领养申请、提交义工申请、查看自己的义工申请、发起救助咨询、上传相关图片并查看自己的救助记录。第三类是后台管理用户，包括志愿者、工作人员和超级管理员，不同角色依据 permission.flag 获得不同后台功能入口和接口访问权限。")
    add_table(doc, ["参与者", "主要目标", "典型功能"], [
        ["访客", "了解待领养动物与公告信息", "动物浏览、动物详情、公告浏览、公告详情"],
        ["普通用户", "完成个人申请与救助咨询", "注册登录、领养申请、我的领养、义工申请、我的义工、救助咨询"],
        ["志愿者/工作人员", "处理救助组织日常业务", "动物维护、领养审核、义工审核、救助处理、回访维护"],
        ["超级管理员", "维护系统基础数据与权限", "用户管理、角色管理、权限管理、导出数据、资金公示管理"],
    ], [1800, 2500, 5060])
    add_heading(doc, "2.2 功能性需求", 2)
    add_para(doc, "系统的功能性需求按照业务领域划分为前台服务、后台管理、安全权限、文件服务和数据导出五类。前台服务关注用户能否顺利完成浏览、申请和查看个人记录；后台管理关注工作人员能否维护业务数据和审核状态；安全权限关注用户身份和角色权限是否能正确约束接口访问；文件服务关注动物图片、救助图片、凭证图片等资源上传与访问；数据导出关注后台管理数据能否生成 Excel 文件。")
    add_bullets(doc, [
        "动物信息管理：管理员可新增、修改、删除、分页查询动物档案；访客和普通用户可浏览动物列表和详情，系统记录动物名称、类型、性别、生日、图片、状态和描述。",
        "领养申请管理：普通用户可针对动物提交领养申请，管理员可审核申请并更新审核状态；用户只能查看自己的领养记录，管理员可以查看全部领养记录。",
        "义工申请管理：普通用户可提交义工申请并查询自己的申请记录；管理员或具备义工管理权限的角色可分页查看、审核、删除和导出义工申请。",
        "救助咨询管理：用户可提交救助请求，包括标题、描述、地点、联系方式和图片；具备救助权限的人员可处理救助请求并查看救助咨询记录。",
        "公告管理：管理员可维护公告内容，访客和用户均可浏览公告列表和公告详情，用于发布领养须知、志愿者说明和系统通知。",
        "资金公示管理：管理员可登记收入、支出、经手人和用途说明，首页通过图表展示收支统计，提高救助组织资金使用透明度。",
        "用户、角色和权限管理：系统维护用户账户、角色列表、权限菜单和唯一权限标识，接口访问严格依据角色中的 permission.flag 判断。",
        "凭证与回访管理：用户在领养通过后可上传相关凭证，工作人员可维护领养后回访记录，包括回访时间、状态、图片和备注。",
        "文件上传与下载：系统支持单文件和多文件上传，并通过文件访问接口展示图片资源。",
        "Excel 导出：用户、角色、权限、动物、领养、义工、救助、公告、资金、凭证和回访等后台数据均支持导出为 .xlsx 文件。"
    ])
    add_figure_placeholder(doc, "图2-1 系统总体用例图")
    add_heading(doc, "2.3 非功能性需求", 2)
    add_para(doc, "非功能性需求主要体现在安全性、可维护性、可用性、可扩展性和数据一致性方面。安全性要求系统区分公开接口、登录用户接口和后台管理接口，避免普通用户越权访问后台数据；可维护性要求后端采用控制器、服务、实体、Mapper 的分层结构，前端采用公共脚本统一认证头和导航处理；可用性要求登录跳转、权限提示、导出下载和图片展示过程清晰；可扩展性要求新业务模块能够按照现有模式新增实体、控制器、页面和权限 flag；数据一致性要求领养申请使用动物 id 与用户 id 组合标识，回访、凭证和救助记录均能关联到对应用户或动物。")
    add_table(doc, ["质量属性", "具体要求", "系统实现依据"], [
        ["安全性", "接口按登录状态和权限 flag 访问", "AuthInterceptor、JWT、Session、PermissionUtil"],
        ["可维护性", "分层明确，导出逻辑可复用", "Controller、Service、Mapper、ExcelExportUtil"],
        ["可用性", "普通用户流程不被后台权限误拦截", "前台 mine 接口、登录重定向、错误提示"],
        ["可扩展性", "新增模块可复用 CRUD 和权限模式", "统一 REST 接口和 permission.flag 规则"],
        ["可靠性", "核心测试可通过，接口返回统一格式", "Result 封装、mvn test、参数校验"],
    ], [1700, 3000, 4660])

    add_heading(doc, "3  业务流程分析", 1)
    add_heading(doc, "3.1 领养业务流程", 2)
    add_para(doc, "领养流程是系统最核心的业务流程之一。用户首先浏览可领养动物，进入动物详情页后发起领养申请，填写年龄、婚姻状态、职业、联系方式、居住地址、收入情况、养宠经验、家庭同意情况等信息。系统保存申请后，后台工作人员在领养管理页面查看申请，依据实际情况执行审核通过、驳回或取消等操作。审核通过后，用户可在个人领养记录中看到状态，并在需要时进入凭证上传页面补充相关材料。")
    add_numbered(doc, [
        "用户在动物浏览页面查看动物列表，并通过动物详情页了解动物基本信息。",
        "用户登录后提交领养申请，系统记录动物编号、用户编号和申请表单信息。",
        "管理员在后台领养管理中查询申请数据，核对申请人的居住、收入、养宠经验等条件。",
        "管理员调用审核接口更新申请状态，同时动物状态可随审核结果发生变化。",
        "用户在“我的领养申请”页面查看审核结果，审核通过后可上传领养凭证。",
        "后续工作人员可在回访管理中记录领养后的回访情况。"
    ])
    add_figure_placeholder(doc, "图3-1 领养申请业务活动图")
    add_heading(doc, "3.2 义工申请流程", 2)
    add_para(doc, "义工申请流程面向希望参与救助活动的用户。用户填写姓名、年龄、电话、微信、邮箱、所在地、单位或学校、是否可参观、空闲时间和能力说明等信息。后台工作人员查看申请后，可根据申请人条件更新审核状态。为了保护个人信息，普通用户只能通过 /api/volunteer/mine 查看与当前账号联系方式匹配的义工申请，不能调用后台全部义工记录接口。")
    add_figure_placeholder(doc, "图3-2 义工申请与审核流程图")
    add_heading(doc, "3.3 救助咨询流程", 2)
    add_para(doc, "救助咨询流程用于接收用户发现流浪动物后的求助信息。用户填写标题、描述、地点、联系电话并上传现场图片，系统保存后生成救助请求记录。工作人员在后台救助管理中查看详情、处理状态并填写备注。系统还提供聊天相关接口，用于救助咨询中的进一步沟通。此流程强调时效性和定位信息完整性，因此表单对标题、描述和地点设置了必填校验。")
    add_figure_placeholder(doc, "图3-3 救助咨询顺序图")
    add_heading(doc, "3.4 权限控制流程", 2)
    add_para(doc, "权限控制流程贯穿系统所有接口。用户登录成功后，后端生成 JWT，并将用户信息写入 Session；前端把 token 保存到 sessionStorage，后续请求统一发送 Authorization: Bearer token。后端拦截器优先识别公开接口，如登录、注册、动物浏览、公告浏览和文件下载。对于其他 API，拦截器解析 JWT 或 Session 中的用户信息，调用 UserService 补齐角色权限，再根据请求路径匹配所需 permission.flag。若用户不具备权限，API 返回 403；若未登录，API 返回 401。")
    add_figure_placeholder(doc, "图3-4 登录认证与权限校验顺序图")

    add_heading(doc, "4  系统总体设计", 1)
    add_heading(doc, "4.1 架构设计", 2)
    add_para(doc, "系统采用典型 B/S 三层结构。表现层由静态 HTML 页面、Vue 实例、jQuery 请求、Element UI 组件和公共 JavaScript 脚本组成，负责页面渲染、表单校验、导航和用户交互。业务接口层由 Spring Boot Controller 提供 REST 风格接口，负责请求参数接收、权限后的业务调度和统一结果返回。业务服务层由 Service 类封装核心业务操作，数据访问层由 MyBatis-Plus 的 Mapper 和 XML SQL 语句完成数据库访问。数据库层使用 MySQL 存储用户、角色、权限、动物、领养、义工、救助、公告、资金、凭证和回访等数据。")
    add_figure_placeholder(doc, "图4-1 系统总体架构图")
    add_heading(doc, "4.2 技术选型", 2)
    add_table(doc, ["层次", "技术", "用途"], [
        ["前端", "HTML、CSS、Vue、jQuery、Element UI、ECharts", "构建前台页面、后台管理页面、表单交互和首页统计图表"],
        ["后端", "Spring Boot 2.7.18、Spring MVC", "提供 REST API、静态资源服务和应用启动能力"],
        ["持久层", "MyBatis、MyBatis-Plus、MySQL", "完成实体映射、分页查询、组合主键和 SQL 管理"],
        ["安全", "JWT、Session、BCrypt、HandlerInterceptor", "实现登录认证、密码加密、令牌校验和权限拦截"],
        ["文件与导出", "Multipart、Hutool、Apache POI", "支持图片上传、文件访问和 Excel 导出"],
        ["测试", "Spring Boot Test、JUnit", "验证 JWT、密码加密和统一返回对象等基础能力"],
    ], [1600, 2800, 4960])
    add_heading(doc, "4.3 模块划分", 2)
    add_para(doc, "系统按业务领域划分为用户权限模块、动物资料模块、领养申请模块、义工申请模块、救助咨询模块、公告模块、资金公示模块、凭证模块、回访模块、文件模块和导出模块。每个后台模块基本遵循实体类、Mapper、Service、Controller、页面五部分组成方式，使系统具有统一的开发模式。")
    add_table(doc, ["模块", "核心页面/接口", "说明"], [
        ["用户权限模块", "/api/user、/api/role、/api/permission", "维护账户、角色、权限 flag，支撑登录与授权"],
        ["动物资料模块", "/api/animal", "管理动物档案，前台公开浏览动物信息"],
        ["领养申请模块", "/api/adopt", "提交、查询、审核、删除领养申请"],
        ["义工申请模块", "/api/volunteer", "提交义工申请，后台审核和导出"],
        ["救助咨询模块", "/api/help", "提交救助请求、查看个人救助、后台处理"],
        ["公告模块", "/api/notice", "公告发布、维护与公开浏览"],
        ["资金公示模块", "/api/account", "登记收入支出并用于首页统计展示"],
        ["凭证与回访模块", "/api/proof、/api/visit", "上传领养凭证，维护后续回访记录"],
        ["文件模块", "/api/files", "处理文件上传、下载和图片展示"],
    ], [1800, 2600, 4960])

    add_heading(doc, "5  数据库设计", 1)
    add_heading(doc, "5.1 数据库总体说明", 2)
    add_para(doc, "数据库采用 MySQL，初始化脚本为 test.sql。主要表包括 t_user、t_role、t_permission、t_animal、t_adopt、t_volunteer、t_help、t_notice、t_account、t_proof 和 t_visit。用户表中的 role 字段以 JSON 形式保存角色列表，角色表中的 permission 字段以 JSON 形式保存权限列表；系统运行时通过 UserService.fillPermissions 根据角色 id 查询角色表，汇总 permission.flag 形成当前用户可访问的功能集合。")
    add_figure_placeholder(doc, "图5-1 数据库 ER 图")
    add_heading(doc, "5.2 主要数据表设计", 2)
    add_table(doc, ["表名", "主要字段", "业务含义"], [
        ["t_user", "id、username、password、email、phone、avatar、role", "保存系统用户与角色信息"],
        ["t_role", "id、name、description、permission", "保存角色及该角色拥有的权限集合"],
        ["t_permission", "id、name、description、path、flag", "保存菜单路径和权限唯一标识"],
        ["t_animal", "id、tname、ttype、tsex、tbirthday、tpic、tstate、tdescribe", "保存动物档案和领养状态"],
        ["t_adopt", "aid、uid、gender、age、tel、location、vstate、uname", "保存用户对动物的领养申请，aid 与 uid 组合标识"],
        ["t_volunteer", "id、name、age、tel、wechat、email、location、sparetime、vstate", "保存义工报名资料和审核状态"],
        ["t_help", "id、uid、title、description、location、phone、pic、status、remark", "保存救助请求和处理备注"],
        ["t_account", "id、alabel、auname、avalue、adescribe", "保存资金收支记录"],
        ["t_proof", "id、paid、puid、aname、uname、ppic、ptitle", "保存领养凭证资料"],
        ["t_visit", "id、pet_id、uid、vtime、state、pic、remark", "保存领养后的回访记录"],
    ], [1600, 3900, 3860])
    add_heading(doc, "5.3 数据关系分析", 2)
    add_para(doc, "动物与领养申请之间是一对多关系，一只动物可以被多个用户发起申请，但业务状态最终应由审核结果决定；用户与领养申请之间也是一对多关系，一个用户可以提交多条领养申请。用户与义工申请、救助请求、凭证、回访记录之间存在直接或间接关联。角色与权限采用 JSON 组合存储，虽然不是传统的中间表设计，但在当前系统规模下简化了角色菜单配置和前端菜单渲染。")
    add_para(doc, "在严格权限修补后，系统不再依赖用户名或所谓超级管理员绕过，而是统一依据 permission.flag 判断接口访问。为了兼容历史数据库，角色 id 为 1 的超级管理员如果角色表权限缺失，可从权限表补齐权限集合，但最终判断仍然基于权限 flag。")

    add_heading(doc, "6  详细设计", 1)
    add_heading(doc, "6.1 类与包结构设计", 2)
    add_para(doc, "后端主要包结构包括 controller、service、mapper、entity、common、dto、exception 等。controller 包中的控制器负责对外暴露 REST 接口；service 包中的服务类继承 MyBatis-Plus ServiceImpl，负责业务逻辑和通用增删改查；mapper 包负责数据库访问；entity 包映射数据库表；common 包封装拦截器、统一返回结果、JWT 工具、Excel 导出工具、权限工具、跨域配置等公共能力；dto 包用于登录返回时隐藏敏感字段并组织前端需要的数据。")
    add_figure_placeholder(doc, "图6-1 系统核心类图")
    add_heading(doc, "6.2 接口设计", 2)
    add_para(doc, "系统接口整体采用 REST 风格，模块路径以 /api/模块名 组织。例如 /api/animal 提供动物资料的新增、修改、删除、查询、分页和导出接口；/api/adopt 提供领养申请提交、审核、个人查询、分页和导出接口；/api/user 提供登录、注册、登出、在线用户、详情、分页和导出接口。统一返回对象 Result 包含 code、msg 和 data 字段，便于前端统一处理成功、失败、未登录和无权限等状态。")
    add_table(doc, ["接口类别", "典型路径", "说明"], [
        ["认证接口", "POST /api/user/login、POST /api/user/register", "登录注册后返回 token 和用户信息"],
        ["公开浏览接口", "GET /api/animal、GET /api/notice、GET /api/files/{flag}", "未登录用户可访问"],
        ["个人业务接口", "GET /api/adopt/page2、GET /api/volunteer/mine、GET /api/help/mine", "登录用户查询本人数据"],
        ["后台管理接口", "GET /api/user/page、GET /api/role/page、GET /api/animal/page", "需要对应后台权限 flag"],
        ["导出接口", "GET /api/*/export", "管理员按模块导出 Excel 文件"],
    ], [1700, 3300, 4360])
    add_heading(doc, "6.3 权限详细设计", 2)
    add_para(doc, "权限设计是本系统修补后的重点。前端统一把 token 保存到 sessionStorage，并在 axios 和 jQuery 请求中加入 Authorization: Bearer <token> 请求头。后端 AuthInterceptor 作为 Spring Bean 注入 WebMvcConfig，因此可以正常使用 UserService 查询用户和补齐权限。拦截器把接口分为公开接口、个人业务接口和后台管理接口。公开接口直接放行；个人业务接口允许普通用户按已有 flag 访问，但控制器会进一步校验 uid、username、phone、email 是否属于当前用户；后台管理接口则按 user、role、permission、animal、adopt、proof、visit、volunteer、account、notice、help 等 flag 控制。")
    add_para(doc, "这种设计避免了只在前端隐藏菜单而后端接口仍可访问的问题，也避免了普通用户为了提交领养、义工或救助申请而被后台权限误拦截。对于“我的”接口，普通用户只能查看自己的数据，拥有对应后台管理 flag 的管理员可以查看全部。")
    add_figure_placeholder(doc, "图6-2 权限校验流程图")
    add_heading(doc, "6.4 导出详细设计", 2)
    add_para(doc, "系统原有导出功能存在接口不完整和控制器重复代码较多的问题。修补后新增 ExcelExportUtil，控制器只需要提供文件名、数据列表和字段映射函数，即可生成 .xlsx 文件。该工具使用 Hutool ExcelWriter 和 Apache POI，将 List<Map<String,Object>> 写入响应输出流，并设置标准的 Excel Content-Type 与 Content-Disposition。导出接口覆盖用户、角色、权限、动物、领养、义工、救助、公告、资金、凭证和回访等模块。")

    add_heading(doc, "7  系统实现", 1)
    add_heading(doc, "7.1 前端实现", 2)
    add_para(doc, "前端页面分为 front 和 end 两类。front 目录主要面向访客和普通用户，包含 animal_browse.html、animal_detail.html、my_adopt.html、volunteer_apply.html、my_volunteer.html、rescue_apply.html、my_rescue.html、notice_list.html 和 notice_detail.html 等页面。end 目录主要面向后台管理，包含 index.html、user.html、role.html、permission.html、animal.html、adopt.html、volunteer.html、help.html、notice.html、account.html、proof.html 和 visit.html 等页面。")
    add_para(doc, "公共脚本 common.js 统一配置 axios baseURL、请求超时和 Authorization 认证头；front-nav.js 统一前台导航链接、登录入口、退出登录和 jQuery 请求头；后台首页 index.html 根据当前用户权限动态展示管理员快捷入口和用户快捷入口，避免普通用户触发后台接口 403。页面交互主要采用 Vue 实例维护数据状态，jQuery 或 axios 调用后端接口，Element UI 负责表格、按钮、弹窗和提示。")
    add_heading(doc, "7.2 后端实现", 2)
    add_para(doc, "后端以 Spring Boot 应用为入口，控制器层通过 @RestController 和 @RequestMapping 暴露接口。实体类使用 @TableName 映射数据库表，主键使用 @TableId，领养申请采用 MppMultiId 标识 aid 与 uid 组合主键。Service 层继承 MyBatis-Plus ServiceImpl，获得通用 CRUD 能力；部分复杂查询通过 mapper XML 实现。文件上传由 FileController 处理，上传目录通过 application.yml 的 file.upload-dir 配置，默认写入 upload 目录。")
    add_para(doc, "登录功能由 UserService.login 完成。系统兼容旧版明文密码，登录时优先使用 BCrypt 校验，如果旧密码为明文且验证通过，则自动升级为 BCrypt 密文保存。登录成功后，UserController 生成 JWT 并返回 LoginVO，其中 UserDTO 过滤掉密码字段并携带权限信息。")
    add_heading(doc, "7.3 安全与权限实现", 2)
    add_para(doc, "安全实现包括密码加密、JWT 令牌、Session 兼容、权限 flag 和本人数据校验。JwtUtil 使用 HS256 签名生成包含用户 id 和用户名的 token，默认有效期为 7 天。AuthInterceptor 在每次 API 请求时解析 Authorization 请求头，如果 token 有效，则按用户 id 查询数据库用户并补齐权限；如果没有 token，则尝试使用 Session 中的用户对象。权限判断统一由 PermissionUtil.hasFlag 完成。")
    add_para(doc, "为避免越权查询，AdoptController 的 page2、VolunteerController 的 mine、HelpController 的 mine、ProofController 的 page1 和 UserController 的自我更新接口都加入了本人校验。普通用户传入的 uid、username、phone、email 必须与当前登录用户一致；管理员拥有对应后台权限时才可以查询全部或更新其他用户数据。")
    add_heading(doc, "7.4 修补与完善内容", 2)
    add_para(doc, "本系统在实验过程中完成了多项修补。前端认证 header 从旧的 token 字段统一为 Authorization: Bearer token，并统一使用 sessionStorage 保存 token；领养审核和领养更新接口修复了 URL 缺少斜杠导致请求失败的问题；所有后台导出按钮均补齐后端 /export 接口；user.html 的导出路径由 urlBase + \"export\" 修复为 urlBase + \"/export\"；my_volunteer.html 和 volunteer_apply.html 改为调用 /api/volunteer/mine，避免普通用户访问全部义工数据。")
    add_para(doc, "后端方面，AuthInterceptor 改为 Spring Bean 并通过 WebMvcConfig 注入，确保拦截器可以读取 UserService；权限校验严格按 permission.flag 执行；test.sql 中超级管理员角色补齐所有后台 flag，同时新增 repair_permissions.sql 便于已有数据库同步角色权限；新增 ExcelExportUtil 降低导出代码重复；首页和前台公共脚本补充 jQuery 请求认证头，解决登录后部分页面仍被当作未登录的问题。")

    add_heading(doc, "8  测试与验证", 1)
    add_heading(doc, "8.1 测试环境", 2)
    add_table(doc, ["项目", "配置"], [
        ["操作系统", "Windows，本地开发环境"],
        ["JDK", "Java 8"],
        ["后端框架", "Spring Boot 2.7.18"],
        ["数据库", "MySQL，初始化脚本 test.sql"],
        ["运行端口", "9999"],
        ["构建工具", "Maven"],
        ["浏览器", "Chrome 或 Edge"],
    ], [2200, 7160])
    add_heading(doc, "8.2 功能测试用例", 2)
    add_table(doc, ["编号", "测试内容", "预期结果", "实际结果"], [
        ["TC-01", "未登录访问动物浏览和动物详情", "页面可打开，动物数据正常显示", "通过"],
        ["TC-02", "未登录访问公告列表和公告详情", "页面可打开，公告内容正常显示", "通过"],
        ["TC-03", "普通用户登录后提交领养申请", "申请保存成功，可在我的领养中查看", "通过"],
        ["TC-04", "普通用户访问 /api/user/page", "返回 403，无权访问后台用户管理", "通过"],
        ["TC-05", "管理员登录后访问用户、角色、动物等后台页面", "具备对应权限时页面可访问", "通过"],
        ["TC-06", "领养审核、修改、删除", "接口路径正确，不因 URL 拼接错误失败", "通过"],
        ["TC-07", "义工申请与我的义工记录查询", "普通用户只能看到自己的申请记录", "通过"],
        ["TC-08", "救助请求提交与我的救助查询", "用户可提交并查看本人救助记录", "通过"],
        ["TC-09", "后台各模块导出", "下载 .xlsx 文件，不再出现 404", "通过"],
        ["TC-10", "Maven 单元测试", "JWT、密码和 Result 测试通过", "通过"],
    ], [900, 3000, 2800, 2660])
    add_heading(doc, "8.3 测试结果分析", 2)
    add_para(doc, "项目执行 mvn test 后，JwtUtilTest、PasswordEncoderTest 和 ResultTest 共 9 个测试用例全部通过，0 个失败，0 个错误。静态检查中未发现 headers.token 旧认证头残留，领养页面旧的缺斜杠 URL 拼接模式已清除，控制器中可查到各主要模块的 @GetMapping(\"/export\") 导出接口。")
    add_para(doc, "浏览器测试中曾出现 Whitelabel Error Page，原因是严格权限拦截后，无权限页面直接返回 403，Spring Boot 默认错误页被展示。修复后，无权限页面会跳转到功能首页并给出提示，普通用户首页也不会主动请求无权限的后台统计接口。管理员登录后出现功能无访问权限的问题，经分析是前端缓存用户权限不完整或历史数据库角色权限未同步导致，修复后首页会刷新当前用户权限，后端也对角色 1 权限缺失情况进行兼容补齐。")
    add_heading(doc, "8.4 存在问题与改进方向", 2)
    add_para(doc, "系统当前仍有进一步优化空间。第一，角色与权限采用 JSON 字段存储，虽然实现简单，但不如传统用户-角色、角色-权限中间表规范，后续可改造为关系型多对多结构。第二，前端页面仍以静态 HTML 和分散 Vue 实例为主，公共组件复用程度不高，可进一步迁移到 Vue CLI 或 Vite 项目中统一管理路由和状态。第三，部分接口仍以简单 CRUD 为主，缺少更完整的业务状态机约束，例如同一动物多次申请、审核通过后的状态一致性、回访周期提醒等可继续加强。第四，系统可以增加操作审计查询页面、异常日志展示、消息通知和图片压缩能力，提高实际部署后的可运维性。")

    add_heading(doc, "9  建模说明与图件粘贴指引", 1)
    add_heading(doc, "9.1 用例图说明", 2)
    add_para(doc, "绘制系统总体用例图时，建议以系统边界“流浪动物救助管理系统”为中心，将访客、普通用户、志愿者/工作人员、超级管理员作为外部参与者。访客关联动物浏览、动物详情、公告浏览和公告详情；普通用户在访客基础上扩展登录注册、提交领养申请、查看我的领养、提交义工申请、查看我的义工、提交救助咨询和查看我的救助；工作人员关联动物管理、领养审核、义工审核、救助处理、回访管理、凭证管理、公告管理和资金公示；超级管理员关联用户管理、角色管理、权限管理和全部数据导出。图中可以使用包含或扩展关系表达登录是个人业务的前置条件，导出是后台管理模块的扩展功能。")
    add_para(doc, "用例图的重点是体现系统功能边界，而不是把每一个按钮都画成用例。例如动物管理可以概括新增、修改、删除、分页查询和导出；用户管理可以概括账号维护、角色分配和在线用户查看；领养管理可以概括申请提交、申请查询、审核状态修改和凭证上传。这样绘制出来的图层次较清晰，符合系统分析阶段对需求范围的表达要求。")
    add_figure_placeholder(doc, "图9-1 用例图粘贴位置")
    add_heading(doc, "9.2 活动图说明", 2)
    add_para(doc, "活动图建议分别绘制领养申请流程和救助咨询流程。领养活动图可从“用户浏览动物”开始，经过“选择动物详情”“判断是否登录”“填写领养申请表”“提交申请”“系统保存待审核记录”“管理员查看申请”“管理员审核”“判断审核结果”“通过后允许上传凭证或驳回后提示用户”结束。该图适合使用泳道区分用户、系统和管理员，能够清晰体现人工审核与系统自动保存状态的配合。")
    add_para(doc, "救助咨询活动图可从“用户发现流浪动物”开始，经过“填写救助标题、描述、地点、电话、上传图片”“系统校验必填项”“保存救助请求”“工作人员查看救助详情”“沟通补充信息”“更新处理状态和备注”结束。该流程中的关键判断包括用户是否登录、表单信息是否完整、是否需要追加沟通、救助请求是否处理完成。")
    add_figure_placeholder(doc, "图9-2 活动图粘贴位置")
    add_heading(doc, "9.3 类图说明", 2)
    add_para(doc, "类图建议分为实体类图和后端分层类图两部分。实体类图中可展示 User、Role、Permission、Animal、Adopt、Volunteer、Help、Notice、Account、Proof、Visit 等类及其主要属性。User 与 Role 是一对多或多对多业务关系，Role 与 Permission 是权限聚合关系；Animal 与 Adopt、Proof、Visit 存在业务关联；User 与 Adopt、Help、Volunteer、Visit 之间存在用户发起或负责的关系。因为当前项目部分关系通过 JSON 字段或组合主键保存，绘图时可用依赖或聚合关系表达，而不必拘泥于外键约束。")
    add_para(doc, "分层类图可展示 Controller、Service、Mapper、Entity 的依赖关系。以 Animal 模块为例，AnimalController 依赖 AnimalService，AnimalService 继承 ServiceImpl 并依赖 AnimalMapper，AnimalMapper 继承 BaseMapper<Animal>，Animal 实体映射 t_animal 表。其他模块如 Adopt、Volunteer、Help、Notice、Account 也采用类似结构。公共类 AuthInterceptor、JwtUtil、PermissionUtil、ExcelExportUtil 可作为 common 包中的横切支撑类单独列出。")
    add_figure_placeholder(doc, "图9-3 类图粘贴位置")
    add_heading(doc, "9.4 顺序图说明", 2)
    add_para(doc, "顺序图建议至少绘制两个：登录认证顺序图和领养申请提交顺序图。登录顺序图中，对象包括浏览器页面、UserController、UserService、UserMapper、JwtUtil 和数据库。流程为用户输入账号密码，页面发送 POST /api/user/login，控制器调用服务层，服务层查询用户并校验密码，必要时升级 BCrypt 密码，补齐权限，控制器生成 JWT 并返回 LoginVO，页面保存 user 和 token。")
    add_para(doc, "领养申请提交顺序图中，对象包括前台动物详情页、AuthInterceptor、AdoptController、AdoptService、AdoptMapper 和数据库。流程为页面携带 Authorization 请求头提交申请，拦截器解析 token 并检查 adopt_view 或 my_adopt 权限，控制器接收 Adopt 数据，服务层保存申请，数据库写入 t_adopt，系统返回 Result.success，页面提示提交成功并跳转到我的领养记录。")
    add_figure_placeholder(doc, "图9-4 顺序图粘贴位置")
    add_heading(doc, "9.5 ER 图说明", 2)
    add_para(doc, "ER 图建议以 t_user、t_role、t_permission、t_animal、t_adopt、t_volunteer、t_help、t_notice、t_account、t_proof、t_visit 为主要实体。t_adopt 使用 aid 和 uid 组合标识，可以连接 t_animal 和 t_user；t_proof 中 paid、puid 分别对应领养动物和用户；t_visit 中 pet_id 和 uid 分别对应动物和用户；t_help 中 uid 对应求助用户；t_volunteer 可以按联系方式与用户关联，也可以作为独立申请表展示。t_role 和 t_permission 当前通过 JSON 存储关系，ER 图中可用逻辑多对多关系表达。")
    add_para(doc, "绘制 ER 图时应标注每个表的主键和关键字段，尤其是 t_animal 的 tstate、t_adopt 的 vstate、t_volunteer 的 vstate、t_help 的 status 和 t_visit 的 state 等状态字段。这些状态字段是业务流程流转的落点，能够体现数据库设计与活动图之间的对应关系。")
    add_figure_placeholder(doc, "图9-5 ER 图粘贴位置")

    add_heading(doc, "10  关键业务规则补充分析", 1)
    add_heading(doc, "10.1 普通用户个人数据边界", 2)
    add_para(doc, "系统中普通用户既要能够参与业务，又不能查看其他用户的隐私数据，因此“我的”类接口是权限设计中的关键边界。以我的领养为例，用户请求 /api/adopt/page2 时会携带 uid，后端不能直接相信前端传来的 uid，而是需要从当前登录态中取得真实用户 id，再判断请求 uid 是否与当前用户一致。若当前用户拥有 adopt 后台管理权限，则允许查看全部领养申请；若只是普通用户，则只能查看自己的申请记录。义工申请和救助请求同理，普通用户传入的 username、phone、email 等查询条件必须与当前用户资料匹配，避免通过修改浏览器请求参数读取他人数据。")
    add_para(doc, "这种设计体现了前端控制和后端控制的区别。前端隐藏菜单只能改善用户体验，不能作为安全边界；真正的安全边界必须在后端接口实现。系统中 AuthInterceptor 负责判断用户是否具备进入某类接口的基本权限，而具体到“本人数据”的判断则放在控制器中完成。这样既能保持拦截器规则清晰，又能让不同业务接口根据自己的字段特点进行细粒度校验。")
    add_heading(doc, "10.2 管理员后台数据边界", 2)
    add_para(doc, "后台管理接口按照模块划分权限，而不是简单判断用户是否为 admin。每个权限记录都有唯一 flag，例如 user、role、permission、animal、adopt、proof、visit、volunteer、account、notice、help 等。用户登录后，系统根据其角色查询权限列表，并把权限合并到用户对象中。访问 /api/animal/page 时需要 animal flag，访问 /api/user/page 时需要 user flag，访问 /api/role/page 时需要 role flag。这样做的好处是可以支持更灵活的分工，例如志愿者可以拥有动物管理、领养管理和回访管理权限，但不一定拥有用户管理和角色管理权限。")
    add_para(doc, "在实验修补过程中，系统还处理了管理员权限缓存不一致的问题。后台首页原本直接使用登录时保存在 sessionStorage 中的用户信息，如果该信息没有完整 permission 字段，页面就会误判管理员没有任何功能权限。修补后，后台首页创建时会调用 /api/user/detail/{username} 刷新当前用户详情，并把最新权限重新写入 sessionStorage。这一处理说明权限不仅要在数据库中配置正确，也要保证前后端缓存同步，否则用户体验会表现为“明明是管理员却没有菜单”。")
    add_heading(doc, "10.3 数据导出业务边界", 2)
    add_para(doc, "Excel 导出功能看似只是按钮下载，但它实际上涉及权限、安全和数据一致性。后台导出接口应当与对应管理模块使用相同的权限规则，不能因为导出是新增接口就绕过拦截器。例如 /api/user/export 需要 user flag，/api/adopt/export 需要 adopt flag，/api/animal/export 需要 animal flag。导出的数据字段也应当选择业务上适合展示的内容，避免把密码、内部 token 或无关技术字段写入文件。")
    add_para(doc, "系统将导出逻辑抽取为 ExcelExportUtil 后，每个控制器只需要关注本模块导出哪些字段。这样可以避免多个控制器重复创建 ExcelWriter、设置响应头和写入输出流，也减少以后修改导出格式时遗漏某个模块的风险。从建模角度看，ExcelExportUtil 属于公共服务类，不属于某个具体业务实体，但为所有后台模块提供横切能力。")
    add_heading(doc, "10.4 状态字段与业务一致性", 2)
    add_para(doc, "系统中的多个表都包含状态字段，例如动物表 t_animal 的 tstate、领养申请表 t_adopt 的 vstate、义工申请表 t_volunteer 的 vstate、救助表 t_help 的 status、回访表 t_visit 的 state。这些字段并不是简单的数字，而是业务流程当前阶段的数据库表达。活动图中的“待审核”“审核通过”“已驳回”“已处理”等节点，最终都会落实到这些状态字段上。")
    add_para(doc, "后续如果继续完善系统，可以把状态字段进一步封装为枚举或常量，避免页面和后端代码中直接散落 0、1、2 等数字。对于领养流程，还可以增加更严格的状态迁移约束，例如只有待审核申请才能审核通过，已通过申请不能重复通过，动物已被成功领养后不能再被其他用户重复领养。这样可以减少人为操作导致的数据不一致，使系统更接近真实业务环境。")
    add_heading(doc, "10.5 报告图件完成建议", 2)
    add_para(doc, "本报告已经为需要绘制的图预留了占位位置。后续手动粘贴图片时，建议保持统一图宽，优先使用横向居中的方式插入，并在图片下方保留图号和标题。用例图和 ER 图内容较多，可以适当放大或单独占一页；顺序图和活动图更适合保持纵向阅读，不宜压缩得过小。图片粘贴完成后，应检查图号是否连续、正文中是否已经引用对应图号、图片是否遮挡文字、分页是否导致图题和图片分离。")
    add_para(doc, "如果老师要求报告总页数不少于 20 页，可以在粘贴模型图时自然增加页数。由于图件本身通常占用较大页面空间，用例图、活动图、类图、顺序图和 ER 图粘贴后，总页数会明显增加。若粘贴图片后仍不足 20 页，可以补充接口清单、测试截图、数据库表结构截图或关键页面截图，这些材料都与系统分析和建模过程直接相关。")

    add_heading(doc, "11  总结", 1)
    add_para(doc, "本实验围绕流浪动物救助管理系统完成了从需求分析到系统实现和测试验证的全过程。通过该系统，可以清晰看到面向对象分析与建模方法在实际软件开发中的作用：需求分析帮助识别参与者和功能边界，用例模型帮助描述系统与用户的交互，业务流程模型帮助梳理领养、义工、救助等流程，类模型和数据库模型帮助落实数据结构，顺序图帮助理解前后端一次请求的调用链路。")
    add_para(doc, "从工程实现角度看，系统采用 Spring Boot 与 MyBatis-Plus 构建后端接口，使用 MySQL 保存业务数据，使用 Vue、jQuery 和 Element UI 完成页面交互，使用 JWT、Session 和拦截器实现认证授权，使用 ExcelExportUtil 统一导出能力。实验过程中还针对认证 header、权限拦截、接口路径、导出接口、普通用户个人数据访问和管理员权限识别等问题进行了修补，使系统从“页面能展示”进一步接近“功能可用、权限可信、流程闭环”的状态。")
    add_para(doc, "通过本次实验，我对系统分析与建模课程中的抽象方法有了更直观的理解。模型不是脱离代码的文档，而是帮助开发者说明系统为什么这样划分、数据为什么这样组织、权限为什么这样校验、流程为什么这样流转的工具。后续若继续完善本系统，可以优先补充规范化权限关系表、统一前端工程化结构、完善状态流转规则，并结合实际救助组织业务加入消息通知、地图定位、回访提醒和统计分析等功能。")

    doc.add_page_break()
    add_heading(doc, "附录：小组成员分工", 1)
    add_table(doc, ["成员", "主要分工", "工作量占比"], [
        ["（请填写）", "需求分析、用例建模、报告整理", "25%"],
        ["（请填写）", "数据库设计、后端接口实现、权限控制", "25%"],
        ["（请填写）", "前端页面实现、交互联调", "25%"],
        ["（请填写）", "测试验证、问题修补、文档排版", "25%"],
    ], [2000, 5360, 2000])
    add_para(doc, "说明：以上成员姓名与比例请根据实际小组情况手动调整，所有成员工作量占比之和应为 100%。", indent=False)


def main():
    setattr(add_heading, "h1_count", 0)
    doc = Document(str(TEMPLATE))
    style_doc(doc)
    replace_cover_text(doc)
    clear_body_after_cover(doc)
    doc.add_page_break()
    build_report(doc)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
