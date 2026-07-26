from pathlib import Path
import shutil
from docx import Document

SRC = Path(r"D:\2026.6.25\文献检索\科技文献检索与论文写作作业3 - 副本.docx")
OUT = Path(r"D:\2026.4.7\Myproject\stray-animal-main\outputs\科技文献检索与论文写作作业3_差异化修改完成版.docx")

REPLACEMENTS = {
    56: "研究内容与方向总结：从代表性成果可以看出，岳晓月老师的研究主线可概括为“食品风险因子识别—功能敏感材料构建—光学或电化学信号转换—移动终端智能判读”。与只追求检出限的传统分析研究相比，这一方向更强调真实食品基质中的选择性、抗干扰能力和现场可操作性。研究对象横跨亚硝酸盐、抗生素、合成抗氧化剂、孔雀石绿及腐败挥发物，说明其关注点已由单一污染物测定扩展到食品加工、贮藏和流通过程中的多类质量安全风险。纳米酶、MOF、比率荧光、水凝胶和微针阵列分别承担催化放大、特异识别、内标校正、便携载体与无损采样等功能，智能手机和机器学习则把颜色信息转化为可量化结果。由此可延伸出的研究方向包括复杂基质校正、多指标阵列传感、模型跨设备迁移、标准样品验证以及快检结果与监管平台的衔接。",
    58: "1. 对学位论文研究链条的理解",
    59: "学位论文不是若干章节的机械组合，而是一条可以被他人检查和复现的研究证据链。其起点应是明确、可回答的问题，例如“不同终压和降压速率如何共同影响叶菜真空预冷的降温速度与失水率”，而不是宽泛地写“研究蔬菜保鲜”。围绕问题，需要进一步界定研究对象、变量、评价指标和适用范围，使题名、摘要、研究目的与结论保持一致。对于食品质量与安全专业，好的论文既要回应品质或安全风险，也要说明技术方案在生产、冷链或监管场景中的实际价值。",
    60: "文献检索在这一过程中承担“建立坐标系”的作用。检索结果不应只按作者和年份排列，而应围绕关键矛盾分类：已有方法解决了什么问题，采用了哪些材料、模型或实验条件，证据强度如何，还存在哪些无法解释的现象。以叶菜真空冷却为例，可把文献分为实验工艺、热质传递模型、结构变形、多目标优化和机器学习代理模型五类，再比较冷却时间、失重率、温度均匀性及模型误差。这样形成的综述能够自然导出研究空白，并为变量选择与技术路线提供依据。",
    61: "研究实施阶段要把“可重复”落实到细节：记录样品品种与成熟度、初始温度和含水量、设备参数、环境条件、重复次数及异常值处理规则；模型研究还应说明边界条件、参数来源、训练集与测试集划分以及误差指标。结果部分负责客观呈现数据，讨论部分则解释机制、比较前人结论并分析局限，两者不能混写。论文定稿前还应核对图表编号、计量单位、统计方法、引文与参考文献的一致性，并如实报告不支持预期的结果。由“问题—证据—方法—结果—解释—边界”构成的闭环，才是学位论文区别于资料汇编的核心。",
    98: "题目：叶菜真空预冷的热质传递建模与机器学习代理预测研究综述",
    99: "叶菜采后仍保持较强的呼吸与蒸腾作用，若田间热不能及时移除，黄化、萎蔫和微生物繁殖会加快。真空预冷通过降低环境压力，使组织中的部分水分在低温下蒸发并吸收潜热，能够在较短时间内降低产品温度，尤其适合比表面积大、含水率高的叶菜。然而，蒸发冷却也会带来失水，叶片、叶柄和叶脉的结构差异还可能造成温度分布不均。因此，工艺评价不能只看降温速度，而应同时考虑冷却时间、失重率、温度均匀性、外观品质和货架期。",
    100: "早期研究主要依靠温度、质量和腔体压力的实验记录分析工艺规律，随后逐步引入有限元和计算流体力学模型。模型通常基于能量守恒与质量守恒，描述多孔组织中的传热、水分迁移和相变蒸发，并利用实测温度曲线与失重数据进行验证。针对鲜切叶菜的研究表明，叶片与叶柄的几何形态和组织结构不同，各位置的降温曲线也存在差异；把所有部位等效为均匀材料，容易掩盖局部热点。近年的热—水—力耦合模型进一步考虑了水分蒸发引起的收缩，以及结构变形对孔隙率和传质过程的反作用，使模型更接近真实品质变化。",
    101: "数值模型的价值不仅在于解释机理，还在于减少大量重复试验。通过改变终压、降压速率、抽气能力、样品尺寸、孔隙率和初始状态，可在计算环境中比较不同条件对冷却时间、失水和均匀性的影响，再结合响应面法或遗传算法寻找折中方案。2026年关于菜心真空冷却的研究采用“实验—模拟—优化”路线，在控制冷却时间、重量损失和温度均匀性之间进行多目标权衡，说明单一最速降温并不必然对应最佳品质。工艺优化必须设置合理约束，并用独立实验验证预测参数，避免模型在数学上最优、在生产中却不可执行。",
    102: "高精度有限元模型计算耗时较长，难以直接用于设备实时控制。机器学习代理模型提供了一条新的路径：先利用经实验验证的物理模型生成覆盖多种工况的时空数据，再训练模型学习输入参数与温度、水分响应之间的映射。Gao等比较KNN、随机森林、XGBoost、MLP、GRU和LSTM后发现，LSTM能够较好重现真空冷却过程的时间演变，并把单工况预测由分钟级缩短到毫秒级。代理模型并非替代物理机理，而是把物理模型的计算结果压缩为快速近似器；其可信度仍取决于训练样本覆盖范围、物理模型准确性和外部实验验证。",
    103: "未来研究可从三方面推进。第一，建立物理约束与数据驱动融合模型，在损失函数或网络结构中嵌入能量和质量守恒，降低小样本条件下的不合理预测。第二，扩大样本来源，纳入品种、成熟度、初始含水量、装载方式和环境波动，并开展跨批次、跨设备验证，评价模型迁移能力。第三，把模型与传感器、可编程控制器和冷链追溯系统结合，根据实时压力、温度与质量变化动态调整抽气过程。只有同时报告预测精度、适用边界、计算成本和品质收益，才能推动真空预冷由离线工艺设计走向可解释、可校准的智能控制。",
    104: "参考文献：",
    105: "[1] Gao H, Zhu Z W, Sun D W. A machine learning surrogate for finite element modelling of vacuum cooling in leafy vegetables: Real-time prediction of heat and mass transfer[J]. Journal of Food Engineering, 2027, 420: 113165. DOI: 10.1016/j.jfoodeng.2026.113165.",
    106: "[2] Gao H, Zhu Z W, Sun D W. Mathematical modelling of vacuum cooling of leafy vegetables using bidirectional conjugate model combining heat and mass transfer with leaf structural deformation[J]. Applied Thermal Engineering, 2025, 258: 124625. DOI: 10.1016/j.applthermaleng.2024.124625.",
    107: "[3] Song X Y, Liu B L, Jaganathan G K. Mathematical simulation on the surface temperature variation of fresh-cut leafy vegetable during vacuum cooling[J]. International Journal of Refrigeration, 2016, 65: 228-237. DOI: 10.1016/j.ijrefrig.2015.12.009.",
    108: "[4] Zhu Z W, et al. Numerical simulation and experimental study of heat and mass transfer in cylinder-like vegetables during vacuum cooling[J]. Innovative Food Science & Emerging Technologies, 2021, 68: 102607. DOI: 10.1016/j.ifset.2021.102607.",
    109: "[5] Dirapan P, Boonyakiat D, Poonlarp P. Improving shelf life, maintaining quality, and delaying microbial growth of broccoli in supply chain using commercial vacuum cooling and package icing[J]. Horticulturae, 2021, 7(11): 506. DOI: 10.3390/horticulturae7110506.",
}

def replace_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)

OUT.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(SRC, OUT)
doc = Document(OUT)
for index, text in REPLACEMENTS.items():
    replace_paragraph_text(doc.paragraphs[index], text)
doc.save(OUT)
print(OUT)
