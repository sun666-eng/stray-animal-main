from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\2026.4.7\Myproject\stray-animal-main")
DOCX = ROOT / "流浪动物救助管理系统-实验报告.docx"
OUT_DOCX = ROOT / "流浪动物救助管理系统-实验报告-含图1-1.docx"
PNG = ROOT / "figures" / "图1-1-流浪动物救助管理系统应用场景示意图.png"


def font(size, bold=False):
    path = r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"
    return ImageFont.truetype(path, size)


def rounded(draw, box, fill, outline="#2F3A4A", width=2, radius=22):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text_center(draw, box, text, fnt, fill="#1F2937", line_gap=8):
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        b = draw.textbbox((0, 0), line, font=fnt)
        widths.append(b[2] - b[0])
        heights.append(b[3] - b[1])
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = box[1] + ((box[3] - box[1]) - total_h) / 2
    for i, line in enumerate(lines):
        x = box[0] + ((box[2] - box[0]) - widths[i]) / 2
        draw.text((x, y), line, font=fnt, fill=fill)
        y += heights[i] + line_gap


def arrow(draw, start, end, color="#475569", width=3):
    draw.line([start, end], fill=color, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 12
    p1 = (end[0] - size * math.cos(angle - math.pi / 6), end[1] - size * math.sin(angle - math.pi / 6))
    p2 = (end[0] - size * math.cos(angle + math.pi / 6), end[1] - size * math.sin(angle + math.pi / 6))
    draw.polygon([end, p1, p2], fill=color)


def draw_icon_user(draw, cx, cy, color):
    draw.ellipse((cx - 18, cy - 34, cx + 18, cy + 2), fill=color)
    draw.rounded_rectangle((cx - 36, cy + 8, cx + 36, cy + 48), radius=18, fill=color)


def draw_icon_db(draw, cx, cy, color):
    draw.ellipse((cx - 42, cy - 25, cx + 42, cy + 5), fill=color, outline="#334155", width=2)
    draw.rectangle((cx - 42, cy - 10, cx + 42, cy + 45), fill=color, outline="#334155", width=2)
    draw.ellipse((cx - 42, cy + 30, cx + 42, cy + 60), fill=color, outline="#334155", width=2)


def draw_icon_doc(draw, cx, cy, color):
    draw.rounded_rectangle((cx - 36, cy - 42, cx + 36, cy + 42), radius=8, fill=color, outline="#334155", width=2)
    draw.line((cx - 20, cy - 14, cx + 20, cy - 14), fill="#334155", width=3)
    draw.line((cx - 20, cy + 2, cx + 20, cy + 2), fill="#334155", width=3)
    draw.line((cx - 20, cy + 18, cx + 12, cy + 18), fill="#334155", width=3)


def build_png():
    PNG.parent.mkdir(exist_ok=True)
    img = Image.new("RGB", (1800, 1180), "#F8FAFC")
    d = ImageDraw.Draw(img)
    title_font = font(44, True)
    h_font = font(27, True)
    body_font = font(23)
    small_font = font(20)

    d.text((90, 55), "流浪动物救助管理系统应用场景示意图", font=title_font, fill="#0F172A")
    d.text((92, 115), "面向访客、普通用户、工作人员和管理员，支撑动物救助、领养、义工、公告、资金与权限管理。", font=body_font, fill="#475569")

    center = (620, 300, 1180, 720)
    rounded(d, center, "#DBEAFE", "#1D4ED8", 4, 34)
    text_center(d, (center[0], center[1] + 24, center[2], center[1] + 105), "流浪动物救助管理系统", h_font, "#1E3A8A")

    modules = [
        ("动物浏览\n动物档案", 690, 430, "#FFFFFF"),
        ("领养申请\n审核凭证", 910, 430, "#FFFFFF"),
        ("义工申请\n救助咨询", 690, 590, "#FFFFFF"),
        ("后台管理\n权限导出", 910, 590, "#FFFFFF"),
    ]
    for txt, x, y, fill in modules:
        box = (x - 95, y - 50, x + 95, y + 50)
        rounded(d, box, fill, "#60A5FA", 2, 18)
        text_center(d, box, txt, small_font, "#1F2937", 5)

    actors = [
        ("访客", "浏览动物\n查看公告", (120, 270, 430, 455), "#E0F2FE", "#0284C7", (275, 342)),
        ("普通用户", "注册登录\n提交申请\n查看本人记录", (120, 565, 430, 790), "#ECFDF5", "#059669", (275, 645)),
        ("志愿者/工作人员", "维护动物\n审核申请\n处理救助回访", (1370, 270, 1680, 495), "#FFF7ED", "#EA580C", (1525, 350)),
        ("超级管理员", "用户角色权限\n公告资金\n全部数据导出", (1370, 585, 1680, 810), "#FEEFEE", "#DC2626", (1525, 665)),
    ]
    for name, desc, box, fill, icon_color, icon_pos in actors:
        rounded(d, box, fill, icon_color, 3, 26)
        draw_icon_user(d, icon_pos[0], icon_pos[1], icon_color)
        d.text((box[0] + 118, box[1] + 35), name, font=h_font, fill="#111827")
        d.text((box[0] + 118, box[1] + 85), desc, font=body_font, fill="#334155", spacing=6)

    arrow(d, (430, 360), (620, 420), "#0284C7")
    arrow(d, (430, 675), (620, 570), "#059669")
    arrow(d, (1370, 380), (1180, 465), "#EA580C")
    arrow(d, (1370, 695), (1180, 600), "#DC2626")

    supports = [
        ("MySQL 数据库", "用户、角色、权限、动物、领养、义工、救助、公告、资金、凭证、回访", (215, 900, 585, 1080), "#EEF2FF", "#4F46E5", "db"),
        ("文件存储", "动物图片、救助图片、领养凭证、回访图片", (715, 900, 1085, 1080), "#F0FDFA", "#0D9488", "doc"),
        ("Excel 导出", "各后台模块导出 .xlsx，便于统计、归档和上报", (1215, 900, 1585, 1080), "#FFFBEB", "#D97706", "doc"),
    ]
    for title, desc, box, fill, color, kind in supports:
        rounded(d, box, fill, color, 3, 22)
        if kind == "db":
            draw_icon_db(d, box[0] + 70, box[1] + 68, color)
        else:
            draw_icon_doc(d, box[0] + 70, box[1] + 78, color)
        d.text((box[0] + 135, box[1] + 34), title, font=h_font, fill="#111827")
        d.text((box[0] + 135, box[1] + 84), desc, font=small_font, fill="#334155", spacing=5)

    arrow(d, (760, 720), (400, 900), "#64748B", 2)
    arrow(d, (900, 720), (900, 900), "#64748B", 2)
    arrow(d, (1035, 720), (1400, 900), "#64748B", 2)

    img.save(PNG, quality=96)
    print(PNG)


def insert_png():
    doc = Document(str(DOCX))
    target = "【此处粘贴：图1-1 流浪动物救助管理系统应用场景示意图】"
    for i, p in enumerate(doc.paragraphs):
        if target in p.text:
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(PNG), width=Inches(6.4))
            try:
                doc.save(str(DOCX))
            except PermissionError:
                doc.save(str(OUT_DOCX))
            print("inserted", i)
            return
    raise RuntimeError("target placeholder not found")


if __name__ == "__main__":
    build_png()
    insert_png()
