from pathlib import Path
from docx import Document

path = Path(r"D:\2026.6.25\文献检索\食品质量与安全542403040115覃祝(1).docx")
doc = Document(path)
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)}")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"P{i:03d}\t[{p.style.name}]\t{text}")
