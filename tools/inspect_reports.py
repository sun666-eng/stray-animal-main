from pathlib import Path
from docx import Document

FILES = [
    Path(r"D:\2026.6.25\文献检索\科技文献检索与论文写作作业3 - 副本.docx"),
    Path(r"D:\2026.6.25\文献检索\科技文献检索与论文写作作业1.docx"),
]

for path in FILES:
    doc = Document(path)
    print(f"\n===== {path.name} =====")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            print(f"P{i:03d}\t[{p.style.name}]\t{text}")
    for ti, table in enumerate(doc.tables):
        print(f"--- TABLE {ti} rows={len(table.rows)} cols={len(table.columns)} ---")
        for ri, row in enumerate(table.rows):
            vals = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            print(f"T{ti}R{ri:02d}\t" + " | ".join(vals))
